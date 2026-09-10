from io import BytesIO

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.shortcuts import redirect

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from django.contrib.staticfiles import finders

from apps.solicitacoes.models import AnexoOPO, CumprimentoOPO, LogSistema, Solicitacao
from apps.solicitacoes.permissoes import escopo_unidades, pode_ver_documentacao_solicitacao


def _autorizado(request):
    return pode_ver_documentacao_solicitacao(request.user)


def _unidade_titulo(request):
    acesso = getattr(request.user, "acesso_institucional", None)
    unidade = getattr(acesso, "unidade", None)
    if unidade:
        return str(unidade.nome).upper()
    cpr = getattr(acesso, "cpr", None)
    if cpr:
        return str(cpr.nome).upper()
    unidades = list(escopo_unidades(request.user).values_list("nome", flat=True).distinct())
    if len(unidades) == 1:
        return str(unidades[0]).upper()
    return "UNIDADE RESPONSÁVEL"


def _geolocalizacao(evento):
    log = LogSistema.objects.filter(
        solicitacao=evento,
        detalhes__icontains="Coordenadas GPS:"
    ).order_by("-criado_em").first()
    if not log:
        return "Não disponível"
    texto = log.detalhes or ""
    pos = texto.find("Coordenadas GPS:")
    if pos < 0:
        return "Não disponível"
    valor = texto[pos + len("Coordenadas GPS:"):].strip()
    return valor.rstrip(".") or "Não disponível"


def _dados(request):
    eventos = (
        Solicitacao.objects
        .filter(unidade__in=escopo_unidades(request.user))
        .select_related("municipio", "unidade")
        .order_by("data_evento", "hora_inicio", "id")
    )
    dados = []
    for evento in eventos:
        opo = AnexoOPO.objects.filter(solicitacao=evento).exclude(arquivo="").order_by("-criado_em").first()
        if not opo:
            continue
        cumprimento = (
            CumprimentoOPO.objects
            .filter(opo=opo, respondido_em__isnull=False)
            .order_by("-respondido_em")
            .first()
        )
        if not cumprimento:
            continue
        if cumprimento.cumprida is True:
            desfecho = "CUMPRIDA"
        elif (cumprimento.justificativa or "").strip():
            desfecho = "JUSTIFICADA"
        else:
            desfecho = "DESCUMPRIDA"
        dados.append({
            "data": evento.data_evento,
            "local": f"{evento.local}{' - ' + evento.municipio.nome if evento.municipio else ''}",
            "evento": evento.nome_evento or "-",
            "desfecho": desfecho,
            "geolocalizacao": _geolocalizacao(evento),
        })
    return dados


@login_required
def relatorio_cumprimentos_pdf(request):
    if not _autorizado(request):
        messages.error(request, "Você não possui acesso aos relatórios de documentação.")
        return redirect("painel_gestao")
    dados = _dados(request)
    response = HttpResponse(content_type="application/pdf")
    response["Content-Disposition"] = 'attachment; filename="relatorio_cumprimento_opos.pdf"'
    doc = SimpleDocTemplate(
        response, pagesize=landscape(A4),
        rightMargin=10 * mm, leftMargin=10 * mm,
        topMargin=10 * mm, bottomMargin=10 * mm,
    )
    styles = getSampleStyleSheet()
    titulo = ParagraphStyle(
        "RelTitulo", parent=styles["Title"], fontName="Helvetica-Bold",
        fontSize=17, leading=20, alignment=TA_CENTER,
        textColor=colors.HexColor("#4b5563"), spaceAfter=3,
    )
    linha = ParagraphStyle(
        "RelLinha", parent=styles["Normal"], fontName="Helvetica-Bold",
        fontSize=10.5, leading=13, alignment=TA_CENTER,
        textColor=colors.HexColor("#4b5563"), spaceAfter=2,
    )
    unidade = ParagraphStyle(
        "RelUnidade", parent=styles["Normal"], fontName="Helvetica-Bold",
        fontSize=11, leading=13, alignment=TA_CENTER,
        textColor=colors.HexColor("#4b5563"), spaceAfter=3,
    )
    celula = ParagraphStyle("RelCelula", parent=styles["Normal"], fontSize=7.5, leading=9)
    cab = ParagraphStyle(
        "RelCab", parent=celula, fontName="Helvetica-Bold",
        textColor=colors.white, alignment=TA_CENTER,
    )
    story = []
    logo_path = finders.find("logos/logo_pmba.png")
    if logo_path:
        logo = Image(logo_path, width=18 * mm, height=18 * mm)
        logo.hAlign = "CENTER"
        story += [logo, Spacer(1, 1.5 * mm)]
    story += [
        Paragraph("POLÍCIA MILITAR DA BAHIA", titulo),
        Paragraph("COMANDO DE OPERAÇÕES POLICIAIS MILITARES", linha),
        Paragraph(_unidade_titulo(request), unidade),
        Paragraph("RELATÓRIO DE CUMPRIMENTO DAS OPOs", linha),
    ]
    rows = [[
        Paragraph("DATA", cab), Paragraph("LOCAL", cab),
        Paragraph("EVENTO", cab), Paragraph("DESFECHO", cab),
        Paragraph("GEOLOCALIZAÇÃO", cab),
    ]]
    for d in dados:
        rows.append([
            Paragraph(d["data"].strftime("%d/%m/%Y"), celula),
            Paragraph(d["local"], celula),
            Paragraph(d["evento"], celula),
            Paragraph(d["desfecho"], celula),
            Paragraph(d["geolocalizacao"], celula),
        ])
    if len(rows) == 1:
        rows.append([Paragraph("Nenhum atendimento registrado.", celula), "", "", "", ""])
    tabela = Table(rows, repeatRows=1, colWidths=[24 * mm, 70 * mm, 70 * mm, 35 * mm, 78 * mm])
    tabela.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4b5563")),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9ca3af")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f3f4f6")]),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story.append(tabela)
    doc.build(story)
    return response


@login_required
def relatorio_cumprimentos_xls(request):
    if not _autorizado(request):
        messages.error(request, "Você não possui acesso aos relatórios de documentação.")
        return redirect("painel_gestao")
    dados = _dados(request)
    wb = Workbook()
    ws = wb.active
    ws.title = "Cumprimento OPO"
    headers = ["DATA", "LOCAL", "EVENTO", "DESFECHO", "GEOLOCALIZAÇÃO"]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="4B5563")
        c.alignment = Alignment(horizontal="center", vertical="center")
    for d in dados:
        ws.append([d["data"], d["local"], d["evento"], d["desfecho"], d["geolocalizacao"]])
    for row in ws.iter_rows(min_row=2):
        row[0].number_format = "DD/MM/YYYY"
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col, width in zip("ABCDE", [14, 40, 42, 20, 55]):
        ws.column_dimensions[col].width = width
    out = BytesIO()
    wb.save(out)
    out.seek(0)
    response = HttpResponse(
        out.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="relatorio_cumprimento_opos.xlsx"'
    return response
