from io import BytesIO
from datetime import datetime

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.shortcuts import redirect

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from apps.solicitacoes.models import AnexoOPO, CumprimentoOPO, LogSistema, Solicitacao
from apps.solicitacoes.permissoes import escopo_unidades, pode_ver_documentacao_solicitacao


def _autorizado(request):
    return pode_ver_documentacao_solicitacao(request.user)


def _dados(request):
    eventos = Solicitacao.objects.filter(
        unidade__in=escopo_unidades(request.user)
    ).select_related("municipio", "unidade").filter(
        opos__arquivo__isnull=False
    ).distinct().order_by("data_evento", "hora_inicio", "id")
    resultado = []
    for evento in eventos:
        opo = AnexoOPO.objects.filter(solicitacao=evento).exclude(arquivo="").order_by("-criado_em").first()
        if not opo:
            continue
        cumprimento = CumprimentoOPO.objects.filter(opo=opo, respondido_em__isnull=False).select_related("operador").order_by("-respondido_em").first()
        desfecho = "NÃO REGISTRADO"
        geolocalizacao = "Não disponível"
        if cumprimento:
            if cumprimento.cumprida is True:
                desfecho = "CUMPRIDA"
            elif cumprimento.cumprida is False:
                desfecho = "JUSTIFICADA" if (cumprimento.justificativa or "").strip() else "DESCUMPRIDA"
            logs = LogSistema.objects.filter(solicitacao=evento).filter(detalhes__icontains="Coordenadas GPS:").order_by("-criado_em")
            if logs.exists():
                texto = logs.first().detalhes
                pos = texto.find("Coordenadas GPS:")
                geolocalizacao = texto[pos + len("Coordenadas GPS:"):].strip().split(". ")[0].strip(" .")
        resultado.append({
            "data": evento.data_evento,
            "local": f"{evento.local}{' - ' + evento.municipio.nome if evento.municipio else ''}",
            "evento": evento.nome_evento,
            "desfecho": desfecho,
            "geolocalizacao": geolocalizacao,
            "protocolo": evento.protocolo,
        })
    return resultado


@login_required
def relatorio_cumprimentos_pdf(request):
    if not _autorizado(request):
        messages.error(request, "Você não possui acesso aos relatórios de documentação.")
        return redirect("painel_gestao")
    dados = _dados(request)
    response = HttpResponse(content_type="application/pdf")
    response["Content-Disposition"] = 'attachment; filename="relatorio_cumprimentos_opo.pdf"'
    doc = SimpleDocTemplate(response, pagesize=landscape(A4), rightMargin=10*mm, leftMargin=10*mm, topMargin=10*mm, bottomMargin=10*mm)
    styles = getSampleStyleSheet()
    titulo = ParagraphStyle("RT", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=16, alignment=TA_CENTER, spaceAfter=3)
    subtitulo = ParagraphStyle("RS", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=11, alignment=TA_CENTER, spaceAfter=8)
    celula = ParagraphStyle("RC", parent=styles["Normal"], fontSize=7.5, leading=9)
    cab = ParagraphStyle("RH", parent=celula, fontName="Helvetica-Bold", textColor=colors.white, alignment=TA_CENTER)
    story = [Paragraph("POLÍCIA MILITAR DA BAHIA", titulo), Paragraph("RELATÓRIO DE CUMPRIMENTO DAS OPOs", subtitulo), Spacer(1, 2*mm)]
    tabela = [[Paragraph("DATA", cab), Paragraph("LOCAL", cab), Paragraph("EVENTO", cab), Paragraph("DESFECHO", cab), Paragraph("GEOLOCALIZAÇÃO", cab)]]
    for d in dados:
        tabela.append([Paragraph(d["data"].strftime("%d/%m/%Y"), celula), Paragraph(d["local"], celula), Paragraph(d["evento"], celula), Paragraph(d["desfecho"], celula), Paragraph(d["geolocalizacao"], celula)])
    if len(tabela) == 1:
        tabela.append([Paragraph("Nenhum atendimento registrado no âmbito disponível.", celula), "", "", "", ""])
    tab = Table(tabela, colWidths=[25*mm, 70*mm, 72*mm, 35*mm, 75*mm], repeatRows=1)
    tab.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#4E342E")),
        ("TEXTCOLOR", (0,0), (-1,0), colors.white),
        ("GRID", (0,0), (-1,-1), .35, colors.HexColor("#9ca3af")),
        ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0,1), (-1,-1), [colors.white, colors.HexColor("#f3f4f6")]),
        ("LEFTPADDING", (0,0), (-1,-1), 5), ("RIGHTPADDING", (0,0), (-1,-1), 5),
        ("TOPPADDING", (0,0), (-1,-1), 5), ("BOTTOMPADDING", (0,0), (-1,-1), 5),
    ]))
    story.append(tab)
    doc.build(story)
    return response


@login_required
def relatorio_cumprimentos_xls(request):
    if not _autorizado(request):
        messages.error(request, "Você não possui acesso aos relatórios de documentação.")
        return redirect("painel_gestao")
    dados = _dados(request)
    wb = Workbook(); ws = wb.active; ws.title = "Cumprimentos OPO"
    cabecalhos = ["DATA", "LOCAL", "EVENTO", "DESFECHO", "GEOLOCALIZAÇÃO"]
    ws.append(cabecalhos)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="4E342E"); c.alignment = Alignment(horizontal="center", vertical="center")
    for d in dados:
        ws.append([d["data"], d["local"], d["evento"], d["desfecho"], d["geolocalizacao"]])
    for row in ws.iter_rows(min_row=2):
        row[0].number_format = "DD/MM/YYYY"
        for cell in row: cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col, width in zip("ABCDE", [14, 38, 42, 20, 48]): ws.column_dimensions[col].width = width
    out = BytesIO(); wb.save(out); out.seek(0)
    response = HttpResponse(out.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = 'attachment; filename="relatorio_cumprimentos_opo.xlsx"'
    return response
