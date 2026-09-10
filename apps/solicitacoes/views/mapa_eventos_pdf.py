from datetime import datetime
from io import BytesIO
import re

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.shortcuts import redirect
from django.contrib.staticfiles import finders

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from apps.solicitacoes.models import Solicitacao, AnexoOPO, CumprimentoOPO, LogSistema
from apps.solicitacoes.permissoes import pode_ver_mapa_eventos, escopo_unidades, pode_ver_documentacao_solicitacao


def _unidade_executor(request, solicitacao):
    acesso = getattr(request.user, "acesso_institucional", None)
    unidade = getattr(acesso, "unidade", None)
    return unidade or solicitacao.unidade


def _tipo_opo_mapa(solicitacao):
    anexo = AnexoOPO.objects.filter(solicitacao=solicitacao).exclude(arquivo="").order_by("-criado_em").first()
    if not anexo:
        return "ORDINÁRIO"
    return "EXTRAORDINÁRIO" if "EVENTO EXTRA: SIM" in (anexo.descricao or "").upper() else "ORDINÁRIO"


def _unidade_titulo(request, eventos=None):
    acesso = getattr(request.user, "acesso_institucional", None)
    unidade = getattr(acesso, "unidade", None)
    if unidade:
        return str(unidade.nome).upper()
    cpr = getattr(acesso, "cpr", None)
    if cpr:
        return str(cpr.nome).upper()
    nomes = []
    if eventos is not None:
        for evento in eventos:
            if evento.unidade and evento.unidade.nome not in nomes:
                nomes.append(evento.unidade.nome)
    if len(nomes) == 1:
        return str(nomes[0]).upper()
    return " / ".join(nomes).upper() if nomes else "UNIDADE RESPONSÁVEL"


def _geolocalizacao(evento):
    """Retorna somente latitude e longitude, sem o campo de precisão."""
    log = LogSistema.objects.filter(solicitacao=evento, detalhes__icontains="Coordenadas GPS:").order_by("-criado_em").first()
    if not log:
        return "Não disponível"
    texto = log.detalhes or ""
    pos = texto.find("Coordenadas GPS:")
    if pos < 0:
        return "Não disponível"
    trecho = texto[pos + len("Coordenadas GPS:"):]
    latitude = re.search(r"latitude\s*=\s*([-+]?\d+(?:\.\d+)?)", trecho, re.IGNORECASE)
    longitude = re.search(r"longitude\s*=\s*([-+]?\d+(?:\.\d+)?)", trecho, re.IGNORECASE)
    if latitude and longitude:
        return f"latitude={latitude.group(1)}, longitude={longitude.group(1)}"
    return "Não disponível"


def _motivos_nao_cumprimento(evento):
    """Extrai do log os motivos marcados no formulário de NÃO cumprimento."""
    log = (LogSistema.objects.filter(solicitacao=evento, acao__icontains="CUMPRIMENTO OPO")
           .filter(detalhes__icontains="Motivos:")
           .order_by("-criado_em").first())
    if not log:
        return ""
    texto = log.detalhes or ""
    match = re.search(r"Motivos:\s*(.*?)(?:\.\s*(?:Observações:|Arquivo de observações:)|$)", texto, re.IGNORECASE | re.DOTALL)
    if not match:
        return ""
    return " / ".join(parte.strip() for parte in match.group(1).split(";") if parte.strip())


def _dados_relatorio_cumprimento(request):
    eventos = (Solicitacao.objects.filter(unidade__in=escopo_unidades(request.user))
               .select_related("municipio", "unidade")
               .order_by("data_evento", "hora_inicio", "id"))
    dados = []
    for evento in eventos:
        opo = AnexoOPO.objects.filter(solicitacao=evento).exclude(arquivo="").order_by("-criado_em").first()
        if not opo:
            continue
        cumprimento = CumprimentoOPO.objects.filter(opo=opo, respondido_em__isnull=False).order_by("-respondido_em").first()
        if not cumprimento:
            continue
        if cumprimento.cumprida is True:
            desfecho = "CUMPRIDA"
        else:
            motivos = _motivos_nao_cumprimento(evento)
            if (cumprimento.justificativa or "").strip():
                desfecho = f"JUSTIFICADA — {motivos}" if motivos else "JUSTIFICADA"
            else:
                desfecho = f"DESCUMPRIDA — {motivos}" if motivos else "DESCUMPRIDA"
        dados.append({
            "data": evento.data_evento,
            "local": f"{evento.local}{' - ' + evento.municipio.nome if evento.municipio else ''}",
            "evento": evento.nome_evento or "-",
            "desfecho": desfecho,
            "geolocalizacao": _geolocalizacao(evento),
        })
    return dados


def _gerar_relatorio_cumprimento_pdf(request):
    dados = _dados_relatorio_cumprimento(request)
    response = HttpResponse(content_type="application/pdf")
    response["Content-Disposition"] = 'attachment; filename="relatorio_cumprimento_opos.pdf"'
    doc = SimpleDocTemplate(response, pagesize=landscape(A4), rightMargin=10*mm, leftMargin=10*mm, topMargin=10*mm, bottomMargin=10*mm)
    styles = getSampleStyleSheet()
    titulo = ParagraphStyle("RTitulo", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=17, leading=20, alignment=TA_CENTER, textColor=colors.HexColor("#4b5563"), spaceAfter=3)
    linha = ParagraphStyle("RLinha", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=10.5, leading=13, alignment=TA_CENTER, textColor=colors.HexColor("#4b5563"), spaceAfter=2)
    unidade = ParagraphStyle("RUnidade", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=11, leading=13, alignment=TA_CENTER, textColor=colors.HexColor("#4b5563"), spaceAfter=3)
    celula = ParagraphStyle("RCel", parent=styles["Normal"], fontSize=7.5, leading=9)
    cab = ParagraphStyle("RCab", parent=celula, fontName="Helvetica-Bold", textColor=colors.white, alignment=TA_CENTER)
    story = []
    logo_path = finders.find("logos/logo_pmba.png")
    if logo_path:
        logo = Image(logo_path, width=18*mm, height=18*mm); logo.hAlign = "CENTER"; story += [logo, Spacer(1, 1.5*mm)]
    story += [
        Paragraph("POLÍCIA MILITAR DA BAHIA", titulo),
        Paragraph("COMANDO DE OPERAÇÕES POLICIAIS MILITARES", linha),
        Paragraph(_unidade_titulo(request), unidade),
        Paragraph("RELATÓRIO DE CUMPRIMENTO DAS OPOs", linha),
    ]
    rows = [[Paragraph("DATA", cab), Paragraph("LOCAL", cab), Paragraph("EVENTO", cab), Paragraph("DESFECHO", cab), Paragraph("GEOLOCALIZAÇÃO", cab)]]
    for d in dados:
        rows.append([Paragraph(d["data"].strftime("%d/%m/%Y"), celula), Paragraph(d["local"], celula), Paragraph(d["evento"], celula), Paragraph(d["desfecho"], celula), Paragraph(d["geolocalizacao"], celula)])
    if len(rows) == 1:
        rows.append([Paragraph("Nenhum atendimento registrado.", celula), "", "", "", ""])
    tabela = Table(rows, repeatRows=1, colWidths=[24*mm, 70*mm, 70*mm, 35*mm, 78*mm])
    tabela.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#4b5563")),
        ("GRID", (0,0), (-1,-1), .4, colors.HexColor("#9ca3af")),
        ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0,1), (-1,-1), [colors.white, colors.HexColor("#f3f4f6")]),
        ("LEFTPADDING", (0,0), (-1,-1), 4), ("RIGHTPADDING", (0,0), (-1,-1), 4),
        ("TOPPADDING", (0,0), (-1,-1), 5), ("BOTTOMPADDING", (0,0), (-1,-1), 5),
    ]))
    story.append(tabela)
    doc.build(story)
    return response


def _gerar_relatorio_cumprimento_xls(request):
    dados = _dados_relatorio_cumprimento(request)
    wb = Workbook(); ws = wb.active; ws.title = "Cumprimento OPO"
    headers = ["DATA", "LOCAL", "EVENTO", "DESFECHO", "GEOLOCALIZAÇÃO"]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="4B5563"); c.alignment = Alignment(horizontal="center", vertical="center")
    for d in dados:
        ws.append([d["data"], d["local"], d["evento"], d["desfecho"], d["geolocalizacao"]])
    for row in ws.iter_rows(min_row=2):
        row[0].number_format = "DD/MM/YYYY"
        for cell in row: cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    for col, width in zip("ABCDE", [14, 40, 42, 35, 55]): ws.column_dimensions[col].width = width
    out = BytesIO(); wb.save(out); out.seek(0)
    response = HttpResponse(out.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = 'attachment; filename="relatorio_cumprimento_opos.xlsx"'
    return response


@login_required
def gerar_mapa_eventos_pdf_seguro(request):
    if request.GET.get("relatorio") == "cumprimentos":
        if not pode_ver_documentacao_solicitacao(request.user):
            messages.error(request, "Você não possui acesso aos relatórios de documentação.")
            return redirect("painel_gestao")
        if request.GET.get("formato") == "xls":
            return _gerar_relatorio_cumprimento_xls(request)
        return _gerar_relatorio_cumprimento_pdf(request)

    if not pode_ver_mapa_eventos(request.user):
        messages.error(request, "O mapa de eventos está disponível para gestores de CPR e Unidade.")
        return redirect("painel_gestao")

    eventos = (Solicitacao.objects.filter(unidade__in=escopo_unidades(request.user))
               .select_related("municipio", "bairro", "unidade")
               .order_by("data_evento", "hora_inicio"))
    data_inicio = (request.GET.get("data_inicio") or "").strip()
    data_fim = (request.GET.get("data_fim") or "").strip()
    if data_inicio:
        try: eventos = eventos.filter(data_evento__gte=datetime.strptime(data_inicio, "%Y-%m-%d").date())
        except ValueError: data_inicio = ""
    if data_fim:
        try: eventos = eventos.filter(data_evento__lte=datetime.strptime(data_fim, "%Y-%m-%d").date())
        except ValueError: data_fim = ""

    response = HttpResponse(content_type="application/pdf")
    response["Content-Disposition"] = 'inline; filename="mapa_eventos.pdf"'
    doc = SimpleDocTemplate(response, pagesize=landscape(A4), rightMargin=10*mm, leftMargin=10*mm, topMargin=10*mm, bottomMargin=10*mm)
    styles = getSampleStyleSheet()
    titulo = ParagraphStyle("TituloMapa", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=17, leading=20, alignment=TA_CENTER, spaceAfter=3)
    comando = ParagraphStyle("ComandoMapa", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=10.5, leading=13, alignment=TA_CENTER, spaceAfter=2)
    unidade = ParagraphStyle("UnidadeMapa", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=11, leading=13, alignment=TA_CENTER, spaceAfter=3)
    subtitulo = ParagraphStyle("SubtituloMapa", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=11, leading=14, alignment=TA_CENTER, spaceAfter=3)
    periodo = ParagraphStyle("PeriodoMapa", parent=styles["Normal"], fontSize=8, alignment=TA_CENTER, textColor=colors.HexColor("#444444"), spaceAfter=8)
    celula = ParagraphStyle("CelulaMapa", parent=styles["Normal"], fontSize=7, leading=8.2)
    cabecalho = ParagraphStyle("CabecalhoMapa", parent=celula, fontName="Helvetica-Bold", textColor=colors.white, alignment=TA_CENTER)
    story = []
    logo_path = finders.find("logos/logo_pmba.png")
    if logo_path:
        logo = Image(logo_path, width=18*mm, height=18*mm); logo.hAlign = "CENTER"; story += [logo, Spacer(1, 1.5*mm)]
    story.append(Paragraph("POLÍCIA MILITAR DA BAHIA", titulo))
    story.append(Paragraph("COMANDO DE OPERAÇÕES POLICIAIS MILITARES", comando))
    story.append(Paragraph(_unidade_titulo(request, eventos), unidade))
    story.append(Paragraph("MAPA DE EVENTO", subtitulo))
    if data_inicio and data_fim:
        periodo_texto = f"Período: {datetime.strptime(data_inicio, '%Y-%m-%d').strftime('%d/%m/%Y')} a {datetime.strptime(data_fim, '%Y-%m-%d').strftime('%d/%m/%Y')}"
    elif data_inicio:
        periodo_texto = f"A partir de {datetime.strptime(data_inicio, '%Y-%m-%d').strftime('%d/%m/%Y')}"
    elif data_fim:
        periodo_texto = f"Até {datetime.strptime(data_fim, '%Y-%m-%d').strftime('%d/%m/%Y')}"
    else:
        periodo_texto = "Todos os eventos do âmbito institucional"
    story.append(Paragraph(periodo_texto, periodo))
    rows = [[Paragraph("DATA", cabecalho), Paragraph("INÍCIO", cabecalho), Paragraph("FIM", cabecalho), Paragraph("EVENTO", cabecalho), Paragraph("MUNICÍPIO", cabecalho), Paragraph("UNIDADE", cabecalho), Paragraph("REGIME", cabecalho)]]
    for evento in eventos:
        rows.append([
            Paragraph(evento.data_evento.strftime("%d/%m/%Y"), celula),
            Paragraph(evento.hora_inicio.strftime("%H:%M") if evento.hora_inicio else "-", celula),
            Paragraph(evento.hora_fim.strftime("%H:%M") if evento.hora_fim else "-", celula),
            Paragraph(str(evento.nome_evento or "-").upper(), celula),
            Paragraph(str(evento.municipio.nome if evento.municipio else "-").upper(), celula),
            Paragraph(str(evento.unidade.nome if evento.unidade else "-").upper(), celula),
            Paragraph(str(_tipo_opo_mapa(evento)).upper(), celula),
        ])
    tabela = Table(rows, repeatRows=1, colWidths=[24*mm,22*mm,22*mm,58*mm,40*mm,82*mm,36*mm], hAlign="CENTER")
    tabela.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#4b5563")),
        ("GRID", (0,0), (-1,-1), .45, colors.HexColor("#9ca3af")),
        ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
        ("LEFTPADDING", (0,0), (-1,-1), 3), ("RIGHTPADDING", (0,0), (-1,-1), 3),
        ("TOPPADDING", (0,0), (-1,-1), 4), ("BOTTOMPADDING", (0,0), (-1,-1), 4),
    ]))
    story.append(tabela)
    story.append(Spacer(1,4*mm))
    story.append(Paragraph("POLÍCIA MILITAR DA BAHIA - Sistema Integrado de Eventos", periodo))
    doc.build(story)
    return response
